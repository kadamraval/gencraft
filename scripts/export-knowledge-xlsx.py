"""
Writes the standard model database as the Knowledge Book workbook (docs/knowledge-book.xlsx) so it can be read as sheets.
The workbook is made from the pack (src/lib/knowledge/pack.json) when it exists; before that it is the empty structure plus the job
vocabulary. The app never reads the workbook; it reads the pack. The two always say the same thing because one is made from the other.

    python scripts/export-knowledge-xlsx.py
"""
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PACK = os.path.join(ROOT, "src", "lib", "knowledge", "pack.json")
JOBS = os.path.join(ROOT, "src", "lib", "knowledge", "jobs.json")
OUT = os.path.join(ROOT, "docs", "knowledge-book.xlsx")

SRC = [("source_url", "Page or API address the value was read from"), ("read_on", "Date it was read (YYYY-MM-DD)")]

TABS = [
    ("1 Makers", "Who builds models, and what do they make?", [
        ("maker_id", "Short unique id, lower case"), ("maker", "Name as the maker writes it"), ("country", "Country of the company"),
        ("docs_url", "Official docs address"), ("own_api", "yes / no"), ("weights", "open / closed / mixed"),
    ]),
    ("2 Models", "Which models exist, and what are they?", [
        ("canonical_id", "maker/family-version: the same on every provider"), ("maker_id", "Link to tab 1"), ("model", "Model name"),
        ("version", "Version as the maker writes it"), ("gen_type", "image / video / audio"), ("category", "Industry name, e.g. text-to-image"),
        ("jobs", "Job ids this model does (tab 3)"), ("released_on", "YYYY-MM-DD"), ("status", "active / preview / deprecated"),
        ("latest_in_family", "yes if it is the newest stable version of its family"), ("replaces", "canonical_id it replaces"),
        ("retires_on", "YYYY-MM-DD if announced"),
    ]),
    ("3 Jobs and names", "The exact vocabulary the app uses, and where each job is used", [
        ("job_id", "Id the code uses"), ("job_name", "Name shown to the user"), ("kind", "tool (changes a file) / create (makes something)"),
        ("gen_type", "image / video / audio"), ("category", "Industry category"), ("takes_in", "What must go in"), ("may_take", "What may also go in"),
        ("gives", "What comes out"), ("also_known_as", "Other names people search for"), ("used_in", "create / audio / edit / story / agent / library"),
    ]),
    ("4 Capabilities", "What can each model do? (one row per model and job)", [
        ("canonical_id", "Link to tab 2"), ("job_id", "Link to tab 3"), ("supported", "yes / no / unknown"), ("seen_by", "Provider names for the job that confirmed it"),
    ]),
    ("5 Parameters", "What exactly can I send? (one row per line and setting)", [
        ("canonical_id", "Link to tab 2"), ("provider", "Link to tab 6"), ("standard_name", "OpenAI-style name, if there is one"), ("provider_field", "Exact field the provider expects"),
        ("type", "string / integer / number / boolean / array"), ("allowed_values", "List, blank if free"), ("default", "Default"), ("min", "Lowest"), ("max", "Highest"),
        ("required", "yes / no"), ("file_kind", "image / video / audio if it takes a file"), ("tier", "common / unique / advanced / technical"),
    ]),
    ("6 Providers", "How does each provider work?", [
        ("provider", "fal / replicate / openrouter / openai / gemini"), ("auth", "How a key is sent"), ("run_style", "sync / async-poll / webhook"),
        ("file_upload", "How input files are sent"), ("rate_limits", "Requests per minute / concurrency"), ("list_read_by", "How its model list is read"),
        ("price_read_by", "How its prices are read"),
    ] + SRC),
    ("7 Provider lines", "Where does each model run, at what price, and which are the backups?", [
        ("canonical_id", "Link to tab 2"), ("provider", "Link to tab 6"), ("endpoint", "Exact model address on that provider"), ("job_id", "Job this line serves"),
        ("rank_in_job", "1 = first choice for this job on this provider, 2 and up = backups"), ("role", "model / substitute"),
        ("price_unit", "per_image / per_second / per_run / per_1m_tokens / per_megapixel / per_1k_chars"), ("price_amount", "USD"), ("price_rules", "Resolution / audio tiers"),
        ("price_from", "standard / litellm / api / page / maker / billed-by-maker"), ("conflict", "yes if two sources differ by more than 10%"),
        ("status", "active / preview / deprecated"),
    ] + SRC),
    ("8 Safety (later)", "Content filters per model (agreed to come later)", [
        ("canonical_id", "Link to tab 2"), ("provider", "Link to tab 6"), ("filter_control", "none / fixed / switchable"), ("can_disable", "yes / no / unknown"),
    ] + SRC),
    ("9 Sources", "How each fact was read, and when", [
        ("provider", "Provider, litellm or maker"), ("method", "standard / litellm / api / page / maker"), ("url", "Address read"), ("read_on", "Date"), ("pinned_to", "For LiteLLM: the commit"),
    ]),
    ("10 Gaps", "Everything NOT FOUND, with the reason (the work list)", [
        ("canonical_id", "Model"), ("provider", "Provider"), ("field", "job / settings / price / endpoint / backups"), ("reason", "Why it is missing"),
    ]),
    ("11 Completeness", "How complete is this?", [
        ("scope", "Sheet, maker or provider"), ("total", "Fields that exist"), ("filled", "Fields read from a source"), ("score", "filled / total"),
    ]),
]

HEAD = PatternFill("solid", fgColor="1F3A5F")
NOTE = PatternFill("solid", fgColor="EEF2F7")


def rows_for(tab, pack, jobs):
    """The data rows for one tab, from the pack (or the job vocabulary alone before the pack exists)."""
    if tab.startswith("3 "):
        return [[j["id"], j["name"], j["kind"], j["group"], j["category"], j["takesIn"], j["mayTake"] or "", j["gives"], ", ".join(j["alsoKnownAs"]), ", ".join(j["usedIn"])] for j in jobs]
    if not pack:
        return []
    out = []
    if tab.startswith("1 "):
        out = [[m["id"], m["name"], m.get("country", ""), m.get("docsUrl", ""), _yn(m.get("ownApi")), m.get("weights", "")] for m in pack["makers"]]
    elif tab.startswith("2 "):
        out = [[m["id"], m["makerId"], m["name"], m.get("version", ""), m["type"], m["category"], ", ".join(m["jobs"]), m.get("releasedOn", ""), m["status"], _yn(m["latest"]), m.get("replaces", ""), m.get("retiresAt", "")] for m in pack["models"]]
    elif tab.startswith("4 "):
        for m in pack["models"]:
            for j in m["jobs"]:
                seen = sorted({s for l in m["lines"] if l["jobId"] == j for s in l["seenBy"]})
                out.append([m["id"], j, "yes", ", ".join(seen)])
    elif tab.startswith("5 "):
        for m in pack["models"]:
            for l in m["lines"]:
                req = set(l["params"].get("required", []))
                for key, p in l["params"]["properties"].items():
                    slot = p.get("format", "")
                    out.append([m["id"], l["provider"], "", key, p.get("type", ""), ", ".join(str(v) for v in p.get("enum", [])), _s(p.get("default")), _s(p.get("minimum")), _s(p.get("maximum")), _yn(key in req), slot, p.get("tier", "")])
    elif tab.startswith("7 "):
        for m in pack["models"]:
            for l in m["lines"]:
                pr = l["price"]
                out.append([m["id"], l["provider"], l["endpoint"], l["jobId"], l["rank"], l["role"], pr["unit"], _s(pr["amount"]), json.dumps(pr.get("rules", [])) if pr.get("rules") else "", pr["from"], _yn(l.get("conflict")), l["status"], pr.get("sourceUrl", ""), pr["readOn"]])
    elif tab.startswith("9 "):
        out = [[s["provider"], s["method"], s["url"], s["readOn"], s.get("pinnedTo", "")] for s in pack["sources"]]
    elif tab.startswith("10 "):
        out = [[g["modelId"], g["provider"], g["field"], g["reason"]] for g in pack["gaps"]]
    elif tab.startswith("11 "):
        out = [[c["scope"], c["total"], c["filled"], c["score"]] for c in pack["completeness"]]
    return out


def _yn(v):
    return "" if v is None else ("yes" if v else "no")


def _s(v):
    return "" if v is None else v


def main():
    jobs = json.load(open(JOBS, encoding="utf8"))["jobs"]
    pack = json.load(open(PACK, encoding="utf8")) if os.path.exists(PACK) else None
    wb = Workbook()
    ws = wb.active
    ws.title = "0 Read me"
    lines = [
        ("Knowledge Book: the standard model database, as sheets",),
        ("Made from the pack the app reads" if pack else "Empty structure plus the job vocabulary. It fills when the build has run.",),
        ("",),
        ("Rules for every cell",),
        ("1. A value is written only if it was read from a source; the address and date are on the same row.",),
        ("2. If it could not be found, it is listed in Gaps. It is never guessed.",),
        ("3. A model is published only when its job, settings, price and endpoint are found on every provider line.",),
        ("4. Source order: OpenAI standard and LiteLLM, then the provider's API, then the scraper.",),
        ("",),
        ("Tab", "Question it answers"),
    ] + [(n, q) for n, q, _ in TABS]
    for r in lines:
        ws.append(list(r))
    ws["A1"].font = Font(bold=True, size=14)
    ws["A4"].font = Font(bold=True)
    for c in ws[10]:
        c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    for name, _q, cols in TABS:
        s = wb.create_sheet(name)
        s.append([c for c, _ in cols])
        s.append([d for _, d in cols])
        for i in range(1, len(cols) + 1):
            h = s.cell(row=1, column=i)
            h.font = Font(bold=True, color="FFFFFF")
            h.fill = HEAD
            d = s.cell(row=2, column=i)
            d.fill = NOTE
            d.font = Font(italic=True, color="555555")
            d.alignment = Alignment(wrap_text=True, vertical="top")
            s.column_dimensions[get_column_letter(i)].width = 24
        s.row_dimensions[2].height = 48
        s.freeze_panes = "A3"
        for r in rows_for(name, pack, jobs):
            s.append(r)
    wb.save(OUT)
    print("saved", OUT, "with" if pack else "without", "pack")


if __name__ == "__main__":
    sys.exit(main())
